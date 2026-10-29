from pathlib import Path
import json, hashlib, re, datetime
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

TASK = Path(__file__).resolve().parent
PROJECT = TASK.parent.parent
MCP = PROJECT / '01_SOURCES/2026-10-27_mukodesi-leiras-mcp_WIP'
OUT = TASK / '2026-10-27_tomeges-nyomtatas-mukodesi-technikai-leiras_WIP.docx'

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def extract_sources():
    entries = []
    for raw in sorted(MCP.glob('*.response.json')):
        response = json.loads(raw.read_text(encoding='utf-8'))
        payload = json.loads('\n'.join(x['text'] for x in response['content'] if x['type'] == 'text'))
        assert payload['success'] is True and payload['status'] == 200
        name = payload.get('class_name') or payload.get('table_name')
        source = payload.get('source_code') or payload.get('table_data')
        extracted = MCP / (name + ('.abap' if 'class_name' in payload else '.table'))
        extracted.write_text(source, encoding='utf-8', newline='')
        assert extracted.read_text(encoding='utf-8') == source.replace('\r\n', '\n')
        entries.append(dict(object=name, tool='GetClass' if 'class_name' in payload else 'GetTable',
                            parameters={('class_name' if 'class_name' in payload else 'table_name'):name, 'version':'active'},
                            response_status=200,
                            saved_at_utc=datetime.datetime.fromtimestamp(raw.stat().st_mtime, datetime.timezone.utc).isoformat(),
                            raw_file=raw.name, raw_sha256=sha(raw), extracted_file=extracted.name,
                            extracted_sha256=sha(extracted)))
    manifest = dict(date='2026-10-27', purpose='Működési és technikai leírás kiegészítő repository-forrásai',
                    scope='Repository source and definitions only; no business table rows',
                    timestamp_note='saved_at_utc is the local persistence time, not a server revision timestamp', objects=entries)
    (MCP / '2026-10-27_mcp-manifest_WIP.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    assert len(entries) == 8
    return entries

def field(p, instr, cached=''):
    r = p.add_run()
    beg = OxmlElement('w:fldChar'); beg.set(qn('w:fldCharType'), 'begin'); r._r.append(beg)
    code = OxmlElement('w:instrText'); code.set(qn('xml:space'), 'preserve'); code.text = instr; r._r.append(code)
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); r._r.append(sep)
    p.add_run(cached)
    endr = p.add_run(); end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); endr._r.append(end)

def cell_shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), fill); pr.append(shade)

def make_table(doc, rows, chapter):
    cols = len(rows[0])
    assert all(len(r) == cols for r in rows), rows
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if cols == 2:
        widths = [3.4, 13.6]
        if chapter == 19: widths = [1.1, 15.9]
    elif cols == 5: widths = [1.5, 8.0, 2.5, 2.5, 2.0]
    elif chapter == 14: widths = [4.5, 4.8, 7.7]
    elif chapter == 13: widths = [3.4, 4.1, 9.5]
    elif chapter == 19: widths = [1.1, 7.4, 8.5]
    else: widths = [3.3, 6.0, 7.7]
    for col, width in zip(table.columns, widths): col.width = Cm(width)
    for ri, values in enumerate(rows):
        cells = table.rows[0].cells if ri == 0 else table.add_row().cells
        trpr = table.rows[ri]._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for ci, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '75'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.line_spacing = 1.05
            run = p.add_run(value); run.font.size = Pt(9 if chapter != 19 else 8.7)
            if ri == 0:
                run.bold = True; run.font.color.rgb = RGBColor.from_string('FFFFFF'); cell_shade(cell, '24445E')
            elif ri % 2 == 0: cell_shade(cell, 'F0F4F7')
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

entries = extract_sources()
text = (TASK / 'leiras.md').read_text(encoding='utf-8')
mapping = {'M05':'M04','M06':'M05','M07':'M06','M08':'M07','M09':'M08'}
if '| M09 |' in text:
    text = re.sub(r'M0[5-9]', lambda m: mapping[m.group()], text)
    (TASK / 'leiras.md').write_text(text, encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21); sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2); sec.right_margin = Cm(2)
sec.footer_distance = Cm(.8)
styles = doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    styles[name].font.name = 'Arial'
    styles[name]._element.get_or_add_rPr().append(OxmlElement('w:lang'))
    styles[name]._element.rPr.find(qn('w:lang')).set(qn('w:val'), 'hu-HU')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.line_spacing = 1.10
styles['Normal'].paragraph_format.space_after = Pt(7)
styles['Normal'].paragraph_format.widow_control = True
styles['Title'].font.size = Pt(28); styles['Title'].font.color.rgb = RGBColor(0,0,0)
styles['Title'].paragraph_format.space_after = Pt(18)
styles['Subtitle'].font.size = Pt(15); styles['Subtitle'].font.color.rgb = RGBColor.from_string('445566')
styles['Heading 1'].font.size = Pt(19); styles['Heading 1'].font.color.rgb = RGBColor.from_string('24445E')
styles['Heading 1'].paragraph_format.page_break_before = True
styles['Heading 1'].paragraph_format.space_after = Pt(13)
styles['Heading 2'].font.size = Pt(12); styles['Heading 2'].font.color.rgb = RGBColor.from_string('24445E')
styles['Heading 2'].paragraph_format.space_before = Pt(12)
styles['Heading 2'].paragraph_format.space_after = Pt(6)
for sn in ['TOC 1','TOC 2']:
    if sn in styles:
        styles[sn].font.name = 'Arial'; styles[sn].font.size = Pt(10)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Tömeges nyomtatás  |  ')
field(footer, ' PAGE ', '1'); footer.add_run(' / '); field(footer,' NUMPAGES ','1')
for r in footer.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string('667788')

doc.core_properties.title = 'Tömeges nyomtatás működési és technikai leírása'
doc.core_properties.subject = 'tomeges-nyomtatas-endre projekt'
doc.core_properties.author = 'Codex'
doc.core_properties.keywords = 'SAP ABAP Inspire DC QM stock dokumentáció'

doc.add_paragraph('Tömeges nyomtatás', 'Title')
doc.add_paragraph('Működési és technikai leírás', 'Subtitle')
doc.add_paragraph('tomeges nyomtatas endre projekt', 'Subtitle')
p = doc.add_paragraph('2026. október 27.'); p.paragraph_format.space_after = Pt(25)
doc.add_paragraph('A csomagok stockba rendezése, az adatok és minták ellenőrzése, az Inspire felé történő engedélyezés, valamint a dokumentumok és e-mailek státuszkövetése egy összefüggő SAP-folyamatot alkot. Ez a leírás a kezelői lépéseket az ABAP-megvalósítással és az adatkapcsolatokkal együtt mutatja be.')
doc.add_paragraph('A folyamat központi állapotsora a nyitott stocktól a DC és QM jóváhagyáson, majd a lezáráson át a teljes kézbesítésig tart. A sikeres SAP-állapotváltást és a külső szolgáltatás eredményét egyaránt figyelembe kell venni.')
make_table(doc, [
 ['Dokumentum adatai','Érték'],
 ['Célközönség','Üzleti kezelők, folyamatszervezők, ABAP-fejlesztők és üzemeltetők'],
 ['Projektforrás','39 objektum mentése 2026. október 26-áról'],
 ['Kiegészítő forrás','1 osztály és 7 tábladefiníció aktív verziója 2026. október 27-éről'],
 ['Fő témák','Stockképzés, DC, QM, lezárás, státuszok, cockpit, adatmodell, integráció és üzemeltetés'],
 ['Visszakereshetőség','Fejezetenkénti forrásjelek, objektumjegyzék és mentett MCP-válaszok']
], 0)

doc.add_page_break()
p = doc.add_paragraph('Tartalomjegyzék')
p.style = styles['Title']; p.runs[0].font.size = Pt(22)
for number, title in enumerate(re.findall(r'^# (.+)$', text, re.M), 1):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(9)
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),f'chapter_{number}')
    run=OxmlElement('w:r'); rpr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'24445E'); rpr.append(color); run.append(rpr)
    node=OxmlElement('w:t'); node.text=title; run.append(node); link.append(run); p._p.append(link)

lines = text.splitlines(); i=0; chapter=0
while i < len(lines):
    line = lines[i].strip()
    if not line: i+=1; continue
    if line.startswith('# '):
        chapter = int(line[2:].split()[0]); heading=doc.add_heading(line[2:], level=1)
        start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(chapter)); start.set(qn('w:name'),f'chapter_{chapter}')
        end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(chapter)); heading._p.insert(1,start); heading._p.append(end)
        i+=1; continue
    if line.startswith('## '):
        doc.add_heading(line[3:], level=2); i+=1; continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals = [x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[: -]+', x) for x in vals): rows.append(vals)
            i+=1
        make_table(doc, rows, chapter); continue
    if line.startswith('Forrás'):
        p=doc.add_paragraph(line); p.paragraph_format.space_before=Pt(6)
        for r in p.runs: r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string('596777')
    else:
        p=doc.add_paragraph(line)
        if re.match(r'^\d+\. ',line):
            p.paragraph_format.left_indent=Cm(.5); p.paragraph_format.first_line_indent=Cm(-.5)
    i+=1

doc.save(OUT)
source_files = list((PROJECT / '05_FINAL/abap-2026-10-26').glob('*.abap'))
source_files += list((PROJECT / '05_FINAL/abap-2026-10-26').glob('*.table'))
source_files += list((PROJECT / '05_FINAL/abap-2026-10-26').glob('*.structure'))
source_files += list(MCP.glob('*.abap')) + list(MCP.glob('*.table'))
inventory = {str(p.relative_to(PROJECT)):sha(p) for p in sorted(source_files)}
(TASK / '2026-10-27_felhasznalt-forrasok_WIP.json').write_text(json.dumps(inventory, indent=2,ensure_ascii=False),encoding='utf-8')
assert len([p for p in doc.paragraphs if p.style.name=='Heading 1']) == 19
print(json.dumps({'docx':str(OUT), 'chapters':19, 'tables':len(doc.tables), 'new_mcp_objects':len(entries), 'paragraphs':len(doc.paragraphs), 'bytes':OUT.stat().st_size},ensure_ascii=False))
